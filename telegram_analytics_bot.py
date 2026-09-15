import glob, os, gc, threading, logging, re, secrets, hashlib, smtplib
from email.mime.text import MIMEText
import pandas as pd
import numpy as np
from datetime import datetime
from io import BytesIO
from flask import Flask

from telegram import InlineKeyboardButton, InlineKeyboardMarkup, Update
from telegram.ext import (
    ApplicationBuilder, CommandHandler, CallbackQueryHandler, 
    MessageHandler, ContextTypes, filters
)
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.chart.data import CategoryChartData

# Logging setup
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

# Flask Web Server
flask_app = Flask(__name__)
@flask_app.route('/')
@flask_app.route('/health')
def health(): 
    return "Analytics Bot Online", 200

def run_flask():
    port = int(os.environ.get("PORT", 10000))
    flask_app.run(host="0.0.0.0", port=port, debug=False, use_reloader=False)

# ---------------------------------------------------------
# AUTHORIZED USERS DIRECTORY & SECURITY DATABASE
# ---------------------------------------------------------
AUTHORIZED_USERS = {
    # Full Admin Access
    "mis2@frozenbottle.in": {"role": "admin", "allowed_stores": "ALL"},
    "mis3@frozenbottle.in": {"role": "admin", "allowed_stores": "ALL"},
    "faraz@frozenbottle.in": {"role": "admin", "allowed_stores": "ALL"},
    "vivek@frozenbottle.in": {"role": "admin", "allowed_stores": "ALL"},

    # Area Managers & Territory Managers
    "am.chennai@frozenbottle.in": {
        "role": "area_manager",
        "allowed_stores": ["Guduvanchery", "Mudichur", "OMR", "Pallikaranai", "Thoraipakkam", "Urapakkam CK", "Velachery"]
    },
    "am1.chennai@frozenbottle.in": {
        "role": "area_manager",
        "allowed_stores": ["Annanagar", "Besant Nagar", "Express Avenue Mall", "Mogappair", "Nanganallur CK", "Race Course Road", "Valsarvakkam", "Vellore"]
    },
    "am2.pune@frozenbottle.in": {
        "role": "area_manager",
        "allowed_stores": ["Baner Road - Pune", "Hinjewadi", "Hinjewadi Phase 3", "Koregaon Park - Pune", "Sinhagad", "Wagholi - CF CK"]
    },
    "am4.chennai@frozenbottle.in": {
        "role": "area_manager",
        "allowed_stores": ["Alwarpet", "Erode", "Iyyappanthangal - CK", "Kolathur", "Nungambakkam - CK", "Perambur - CK", "Zamin Pallavaram"]
    },
    "areamanager.kerala@frozenbottle.in": {
        "role": "area_manager",
        "allowed_stores": ["Kakkanad", "Ravipuram", "Thiruvalla"]
    },
    "areamanager.mumbai3@frozenbottle.in": {
        "role": "area_manager",
        "allowed_stores": ["Badlapur", "Byculla", "Kalyan", "Khar", "Lokhandwala", "Malad - CF - CK", "Prabhadevi", "Thakur Village"]
    },
    "areamanager.mumbai@frozenbottle.in": {
        "role": "area_manager",
        "allowed_stores": ["Dahisar", "Kamothe", "Manpada - CF CK", "Marol - CF CK", "Mira Road", "Mulund", "Powai- CF - CK", "SEAWOOD", "Sher- E-Punjab", "Virar"]
    },
    "areamanager1@frozenbottle.in": {
        "role": "area_manager",
        "allowed_stores": ["AECS Layout", "Gunjur", "ITPL", "Kempfort", "Manipal", "Miraya Rose", "Shivamogga", "Tata Sherwood", "Tumkur", "Whitefield", "Yemalur"]
    },
    "areamanager5@frozenbottle.in": {
        "role": "area_manager",
        "allowed_stores": ["Ananth Nagar", "BTM Layout", "Banashankari", "HSR Layout", "Harlur Road", "JP Nagar", "Kadubisanahalli - CF CK", "Koramangala", "Meenakshi Mall", "Sarjapur Road"]
    },
    "areamanager6@frozenbottle.in": {
        "role": "area_manager",
        "allowed_stores": ["Basaveshwarnagar", "Bel Road", "Channasandra", "Frazer Town", "Indiranagar - CK", "Kammanhalli", "Kanakapura", "Kolar- Highway Star", "Nagavara", "Yelahanka"]
    },
    "ops@lubov.in": {
        "role": "area_manager",
        "allowed_stores": ["Lubov Store"]
    },
    "bangaloreterritorymanager1@frozenbottle.in": {
        "role": "territory_manager",
        "allowed_stores": ["AECS Layout", "Ananth Nagar", "BTM Layout", "Banashankari", "Basaveshwarnagar", "Bel Road", "Channasandra", "Frazer Town", "Gunjur", "HSR Layout", "Harlur Road", "ITPL", "Indiranagar - CK", "JP Nagar", "Kadubisanahalli - CF CK", "Kammanhalli", "Kanakapura", "Kempfort", "Kolar- Highway Star", "Koramangala", "Lubov Store", "Manipal", "Meenakshi Mall", "Miraya Rose", "Nagavara", "Sarjapur Road", "Shivamogga", "Tata Sherwood", "Tumkur", "Whitefield", "Yelahanka", "Yemalur"]
    },
    "bangaloreterritorymanager@frozenbottle.in": {
        "role": "territory_manager",
        "allowed_stores": ["Kakkanad", "Ravipuram", "Thiruvalla"]
    },
    "manohar@frozenbottle.in": {
        "role": "territory_manager",
        "allowed_stores": ["Badlapur", "Baner Road - Pune", "Byculla", "Dahisar", "Hinjewadi", "Hinjewadi Phase 3", "Kalyan", "Kamothe", "Khar", "Koregaon Park - Pune", "Lokhandwala", "Malad - CF - CK", "Manpada - CF CK", "Marol - CF CK", "Mira Road", "Mulund", "Powai- CF - CK", "Prabhadevi", "SEAWOOD", "Sher- E-Punjab", "Sinhagad", "Thakur Village", "Virar", "Wagholi - CF CK"]
    },
    "tm.chennai@frozenbottle.in": {
        "role": "territory_manager",
        "allowed_stores": ["Alwarpet", "Annanagar", "Besant Nagar", "Erode", "Express Avenue Mall", "Guduvanchery", "Iyyappanthangal - CK", "Kolathur", "Mogappair", "Mudichur", "Nanganallur CK", "Nungambakkam - CK", "OMR", "Pallikaranai", "Perambur - CK", "Race Course Road", "Thoraipakkam", "Urapakkam CK", "Valsarvakkam", "Velachery", "Vellore", "Zamin Pallavaram"]
    }
}

# Dynamic In-Memory User Security Store: email -> {"hash": sha256_str, "plain": plain_text_pass}
USER_PASSWORDS = {}
SESSION_CACHE = {}

def hash_pass(pwd: str) -> str:
    return hashlib.sha256(pwd.encode()).hexdigest()

def generate_random_password(length=8):
    return secrets.token_hex(length // 2)

def send_password_email(to_email, raw_password):
    smtp_server = os.environ.get("SMTP_SERVER")
    smtp_port = os.environ.get("SMTP_PORT", 587)
    smtp_email = os.environ.get("SMTP_EMAIL")
    smtp_password = os.environ.get("SMTP_PASSWORD")

    if not all([smtp_server, smtp_email, smtp_password]):
        logger.warning("SMTP environment variables not configured. Skipping email send.")
        return False

    try:
        msg = MIMEText(
            f"Hello,\n\nYour login password for the Frozen Bottle Analytics Telegram Bot is:\n\nPassword: {raw_password}\n\n"
            f"Please keep this password safe.\n\nRegards,\nAnalytics Team"
        )
        msg['Subject'] = "Your Analytics Bot Access Password"
        msg['From'] = smtp_email
        msg['To'] = to_email

        with smtplib.SMTP(smtp_server, int(smtp_port)) as server:
            server.starttls()
            server.login(smtp_email, smtp_password)
            server.sendmail(smtp_email, [to_email], msg.as_string())
        return True
    except Exception as e:
        logger.error(f"Failed to send email to {to_email}: {e}")
        return False

# ---------------------------------------------------------
# Parquet Data Loader & Memory Optimization
# ---------------------------------------------------------
DB_CACHE_FILE = "cached_dataset.parquet"

def optimize_and_cache_data():
    if os.path.exists(DB_CACHE_FILE):
        return pd.read_parquet(DB_CACHE_FILE)

    all_csvs = sorted(list(set(glob.glob("**/*.csv", recursive=True) + glob.glob("/home/RaviMallappa/**/*.csv", recursive=True))))
    if not all_csvs: 
        raise FileNotFoundError("No historical CSV files found!")

    target_cols = ['Date', 'Brand Name', 'Brand', 'Branch', 'Store', 'Store Type', 'Region', 'Source', 'Session', 'Net Sales', 'Orders', 'Discount', 'Gross Sales']
    
    parquet_parts = []
    for f in all_csvs:
        try:
            s_df = pd.read_csv(f, nrows=1)
            v_cols = [c for c in target_cols if c in s_df.columns]
            df_part = pd.read_csv(f, usecols=v_cols, low_memory=True)

            df_part['Date'] = pd.to_datetime(df_part['Date'], errors='coerce')
            df_part = df_part.dropna(subset=['Date'])
            if df_part.empty:
                continue

            df_part['Brand Name'] = df_part.get('Brand Name', df_part.get('Brand', 'Unknown')).astype(str).astype('category')
            df_part['Branch'] = df_part.get('Branch', df_part.get('Store', 'Unknown')).astype(str).astype('category')
            df_part['Source'] = df_part.get('Source', 'Unknown').astype(str).astype('category')
            if 'Store Type' in df_part: df_part['Store Type'] = df_part['Store Type'].astype(str).astype('category')
            if 'Region' in df_part: df_part['Region'] = df_part['Region'].astype(str).astype('category')
            if 'Session' in df_part: df_part['Session'] = df_part['Session'].astype(str).astype('category')

            for num_col in ['Net Sales', 'Gross Sales', 'Discount']:
                df_part[num_col] = (pd.to_numeric(df_part.get(num_col, 0), errors='coerce').fillna(0.0) / 100000.0).astype('float32')
            df_part['Orders'] = pd.to_numeric(df_part.get('Orders', 0), errors='coerce').fillna(0).astype('int32')

            parquet_parts.append(df_part)
            del df_part
            gc.collect()
        except Exception:
            continue

    if not parquet_parts:
        raise ValueError("No valid data could be processed from CSVs.")

    df = pd.concat(parquet_parts, ignore_index=True)
    del parquet_parts
    gc.collect()

    df['YearMonth'] = df['Date'].dt.strftime('%Y-%m').astype('category')
    df['MonthLabel'] = df['Date'].dt.strftime('%b %Y').astype('category')

    raw_sales_inr = df['Net Sales'] * 100000.0
    df['Calc_AOV'] = np.where(df['Orders'] > 0, raw_sales_inr / df['Orders'], 0.0).astype('float32')
    df['Calc_Disc_Pct'] = np.where(df['Gross Sales'] > 0, (df['Discount'] / df['Gross Sales']) * 100, 0.0).astype('float32')

    df['AOV Bucket'] = pd.cut(df['Calc_AOV'], bins=[-np.inf, 200, 400, 600, 800, np.inf], labels=['< ₹200', '₹200 - ₹400', '₹400 - ₹600', '₹600 - ₹800', '> ₹800'])
    df['Discount Bucket'] = pd.cut(df['Calc_Disc_Pct'], bins=[-np.inf, 5, 15, 25, 35, np.inf], labels=['0 - 5%', '5 - 15%', '15 - 25%', '25 - 35%', '> 35%'])

    df.to_parquet(DB_CACHE_FILE, compression='snappy')
    return df

GLOBAL_DF = optimize_and_cache_data()
DIM_COL_MAP = {
    'Brand': 'Brand Name',
    'Region': 'Region',
    'Source': 'Source',
    'Session': 'Session',
    'Store': 'Branch',
    'AOV Bucket': 'AOV Bucket',
    'Discount Bucket': 'Discount Bucket'
}

def generate_pivot(df_filtered, dim_col, months):
    df_eval = df_filtered[df_filtered['YearMonth'].isin(months)].copy()
    if df_eval.empty or dim_col not in df_eval.columns:
        return pd.DataFrame()
        
    piv = pd.pivot_table(df_eval, index=dim_col, columns='YearMonth', values='Net Sales', aggfunc='sum', fill_value=0)
    
    if len(piv.columns) >= 2:
        c1, c2 = piv.columns[-2], piv.columns[-1]
        piv['MoM Growth %'] = np.where(piv[c1] > 0, ((piv[c2] - piv[c1]) / piv[c1]) * 100, 0.0)
    piv['Total Sales (Lacs)'] = piv[[c for c in piv.columns if c != 'MoM Growth %']].sum(axis=1)
    
    return piv.sort_values(by='Total Sales (Lacs)', ascending=False)

def generate_store_split_pivot(df_filtered, sec_dim_col, months, is_store_prim=False):
    df_eval = df_filtered[df_filtered['YearMonth'].isin(months)].copy()
    if df_eval.empty:
        return pd.DataFrame()
        
    if is_store_prim and sec_dim_col != 'Branch':
        idx_cols = ['Branch', sec_dim_col]
    else:
        idx_cols = [sec_dim_col]

    piv = pd.pivot_table(df_eval, index=idx_cols, columns='YearMonth', values='Net Sales', aggfunc='sum', fill_value=0)
    
    if len(piv.columns) >= 2:
        c1, c2 = piv.columns[-2], piv.columns[-1]
        piv['MoM Growth %'] = np.where(piv[c1] > 0, ((piv[c2] - piv[c1]) / piv[c1]) * 100, 0.0)
        
    piv['Total Sales (Lacs)'] = piv[[c for c in piv.columns if c != 'MoM Growth %']].sum(axis=1)
    return piv.sort_values(by=['Total Sales (Lacs)'], ascending=False)

def get_filtered_data(filters_dict, timeframe, user_config):
    df = GLOBAL_DF.copy()

    # Store scoping per user
    allowed_stores = user_config.get('allowed_stores', 'ALL')
    if allowed_stores != 'ALL':
        df = df[df['Branch'].isin(allowed_stores)]

    if filters_dict.get('Store Type') and filters_dict['Store Type'] != 'ALL':
        df = df[df['Store Type'] == filters_dict['Store Type']]
    for k, col in DIM_COL_MAP.items():
        sel = filters_dict.get(k, set())
        if sel and 'ALL' not in sel and col in df.columns:
            df = df[df[col].isin(list(sel))]
    
    avail = sorted(df['YearMonth'].dropna().unique().tolist())
    if not avail:
        return df, df, []

    # Filter logic to handle current month vs completed historical months
    if timeframe == "Current Month":
        months = [avail[-1]]
    elif timeframe == "Last Month":
        months = [avail[-2]] if len(avail) >= 2 else [avail[-1]]
    else:
        completed_months = avail[:-1] if len(avail) > 1 else avail
        tf_map = {
            "Last 2 Months": 2,
            "Quarterly": 3,
            "Half-Yearly": 6,
            "Yearly": 12
        }
        count = tf_map.get(timeframe, 3)
        months = completed_months[-count:]
    
    return df, df[df['YearMonth'].isin(months)].copy(), months

def build_telegram_summary(piv, primary_dim, timeframe):
    if piv.empty:
        return "No data available for the selected parameters."
    lines = [f"📊 **Performance Summary in Lacs ({primary_dim} | {timeframe})**\n"]
    lines.append("`" + f"{primary_dim[:12]:<12} | " + " | ".join([str(c) for c in piv.columns[:-2]]) + " | Total`")
    lines.append("`" + "-"*40 + "`")
    
    for idx, row in piv.head(8).iterrows():
        val_str = " | ".join([f"₹{v:.2f}L" for v in row[:-2]])
        tot_str = f"₹{row['Total Sales (Lacs)']:.2f}L"
        lines.append(f"`{str(idx)[:12]:<12} | {val_str} | {tot_str}`")
        
    tot_sales = piv['Total Sales (Lacs)'].sum()
    lines.append("\n" + f"💰 **Total Period Net Sales:** ₹{tot_sales:,.2f} Lacs")
    return "\n".join(lines)

def build_multi_sheet_excel(df_filtered, months, primary_dim='Store'):
    out = BytesIO()
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    is_store_prim = (primary_dim == 'Store')

    sections = [
        ("Store Summary", "Branch"),
        ("Brand Summary", "Brand Name"),
        ("Source Summary", "Source"),
        ("Session Summary", "Session"),
        ("AOV Bucket Summary", "AOV Bucket"),
        ("Discount Bucket Summary", "Discount Bucket")
    ]

    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(border_style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for sheet_title, dim_col in sections:
        piv = generate_store_split_pivot(df_filtered, dim_col, months, is_store_prim)
        ws = wb.create_sheet(title=sheet_title)
        
        ws.append([f"{sheet_title} Report (in ₹ Lacs)"])
        ws.cell(1, 1).font = Font(size=14, bold=True, color="1F4E78")
        ws.append([])

        if piv.empty:
            ws.append(["No data available"])
            continue

        reset_piv = piv.reset_index()
        headers = list(reset_piv.columns)
        ws.append(headers)

        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(3, col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        for r in reset_piv.values:
            row_vals = []
            for val in r:
                if isinstance(val, (float, np.floating, int, np.integer)):
                    row_vals.append(round(float(val), 2))
                else:
                    row_vals.append(val)
            ws.append(row_vals)

        sum_row = ["Total Summary (Lacs)"]
        if is_store_prim and dim_col != 'Branch':
            sum_row.append("")

        for c in piv.columns:
            if c == 'MoM Growth %':
                sum_row.append("-")
            else:
                sum_row.append(round(float(piv[c].sum()), 2))
        ws.append(sum_row)

        mom_col_idx = headers.index('MoM Growth %') + 1 if 'MoM Growth %' in headers else None

        for r_idx, row in enumerate(ws.iter_rows(min_row=4, max_row=ws.max_row, min_col=1, max_col=len(headers)), start=4):
            for c_idx, cell in enumerate(row, start=1):
                cell.border = border
                if isinstance(cell.value, (int, float)):
                    if mom_col_idx and c_idx == mom_col_idx:
                        cell.number_format = '0.0"%"'
                    else:
                        cell.number_format = '₹#,##0.00'

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            ws.column_dimensions[get_column_letter(col[0].column)].width = max(max_len + 3, 14)

    ws_raw = wb.create_sheet(title="Raw Data (Sales in Lacs)")
    ws_raw.append(list(df_filtered.columns))
    for r in df_filtered.head(5000).values:
        ws_raw.append([str(x) if isinstance(x, pd.Timestamp) else x for x in r])

    wb.save(out)
    out.seek(0)
    return out

def add_analysis_slide(prs, title, piv, dim_name):
    if piv.empty:
        return
        
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)
    
    tb_title = slide.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(12), Inches(0.6))
    p_title = tb_title.text_frame.paragraphs[0]
    p_title.text = f"{title} (in ₹ Lacs)"
    p_title.font.size = Pt(24)
    p_title.font.bold = True
    p_title.font.color.rgb = RGBColor(31, 78, 120)

    chart_data = CategoryChartData()
    categories = list(piv.head(6).index.astype(str))
    chart_data.categories = categories

    month_cols = [c for c in piv.columns if c not in ['MoM Growth %', 'Total Sales (Lacs)']]
    for m in month_cols:
        series_vals = [round(float(v), 2) for v in piv.head(6)[m]]
        chart_data.add_series(str(m), series_vals)

    x, y, cx, cy = Inches(0.6), Inches(1.2), Inches(7.5), Inches(5.5)
    chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, x, y, cx, cy, chart_data).chart
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.TOP
    chart.plots[0].has_data_labels = True
    
    for series in chart.series:
        for point in series.points:
            dl = point.data_label
            dl.font.size = Pt(8)
            dl.number_format = '0.00'

    tb_insight = slide.shapes.add_textbox(Inches(8.3), Inches(1.2), Inches(4.5), Inches(5.5))
    tf = tb_insight.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "📌 Key Insights & Performance"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = RGBColor(31, 78, 120)
    
    top_performer = piv.index[0]
    top_sales = piv.iloc[0]['Total Sales (Lacs)']
    p1 = tf.add_paragraph()
    p1.text = f"• Top Contributor: {top_performer} with ₹{top_sales:.2f} Lacs net sales."
    p1.font.size = Pt(12)
    
    if 'MoM Growth %' in piv.columns:
        valid_growth = piv.dropna(subset=['MoM Growth %'])
        if not valid_growth.empty:
            highest_growth = valid_growth.sort_values(by='MoM Growth %', ascending=False).iloc[0]
            lowest_growth = valid_growth.sort_values(by='MoM Growth %', ascending=True).iloc[0]
            
            p2 = tf.add_paragraph()
            p2.text = f"• Highest Growth: {highest_growth.name} ({highest_growth['MoM Growth %']:+.1f}% MoM)."
            p2.font.size = Pt(12)
            
            p3 = tf.add_paragraph()
            p3.text = f"• Drop/Lagging Area: {lowest_growth.name} ({lowest_growth['MoM Growth %']:+.1f}% MoM)."
            p3.font.size = Pt(12)

    total_sales = piv['Total Sales (Lacs)'].sum()
    top_3_contrib = (piv.head(3)['Total Sales (Lacs)'].sum() / total_sales * 100) if total_sales > 0 else 0
    p4 = tf.add_paragraph()
    p4.text = f"• Concentration: Top 3 {dim_name}s drive {top_3_contrib:.1f}% of total sales."
    p4.font.size = Pt(12)

def build_pptx(df_filtered, months, primary_dim, timeframe):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.33), Inches(7.5)
    blank_layout = prs.slide_layouts[6]
    
    s1 = prs.slides.add_slide(blank_layout)
    bg1 = s1.shapes.add_shape(1, 0, 0, Inches(13.33), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = RGBColor(31, 78, 120)
    
    tb = s1.shapes.add_textbox(Inches(1), Inches(2.5), Inches(11.33), Inches(2))
    p = tb.text_frame.paragraphs[0]
    p.text = "Executive Performance Overview"
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    
    p2 = tb.text_frame.add_paragraph()
    p2.text = f"Primary Focus: {primary_dim} | Timeframe: {timeframe} | Figures in ₹ Lacs"
    p2.font.size = Pt(20)
    p2.font.color.rgb = RGBColor(200, 220, 240)

    dimensions = [
        ("Brand Breakdown & Insights", "Brand Name", "Brand"),
        ("Source Contribution & Insights", "Source", "Source"),
        ("Session Performance & Insights", "Session", "Session"),
        ("AOV Bucket Distribution", "AOV Bucket", "AOV Bucket"),
        ("Discount Bucket Distribution", "Discount Bucket", "Discount Bucket")
    ]

    for title, dim_col, dim_name in dimensions:
        piv = generate_pivot(df_filtered, dim_col, months)
        add_analysis_slide(prs, title, piv, dim_name)

    out = BytesIO()
    prs.save(out)
    out.seek(0)
    return out

# UI Menus
def get_main_menu():
    kb = [
        [InlineKeyboardButton("🏷️ Brand", callback_data="p_Brand"), InlineKeyboardButton("🏬 Store", callback_data="p_Store")],
        [InlineKeyboardButton("🗺️ Region", callback_data="p_Region"), InlineKeyboardButton("🌐 Source", callback_data="p_Source")],
        [InlineKeyboardButton("🕒 Session", callback_data="p_Session"), InlineKeyboardButton("💰 AOV Bucket", callback_data="p_AOV Bucket")],
        [InlineKeyboardButton("🏷️ Discount Bucket", callback_data="p_Discount Bucket")]
    ]
    return InlineKeyboardMarkup(kb)

def get_store_type_menu():
    kb = [
        [InlineKeyboardButton("🌐 ALL Types", callback_data="st_ALL")],
        [InlineKeyboardButton("🏬 FOFO", callback_data="st_FOFO"), InlineKeyboardButton("🏢 COCO", callback_data="st_COCO")],
        [InlineKeyboardButton("🤝 Partner", callback_data="st_Partner")]
    ]
    return InlineKeyboardMarkup(kb)

def get_timeframe_menu():
    kb = [
        [InlineKeyboardButton("📅 Current Month", callback_data="tf_Current Month"), InlineKeyboardButton("📅 Last Month", callback_data="tf_Last Month")],
        [InlineKeyboardButton("📊 Last 2 Months", callback_data="tf_Last 2 Months"), InlineKeyboardButton("📈 Quarterly", callback_data="tf_Quarterly")],
        [InlineKeyboardButton("📉 Half-Yearly", callback_data="tf_Half-Yearly"), InlineKeyboardButton("📅 Yearly", callback_data="tf_Yearly")]
    ]
    return InlineKeyboardMarkup(kb)

def get_filter_menu(dim_name, selected_set):
    col = DIM_COL_MAP[dim_name]
    opts = sorted(GLOBAL_DF[col].dropna().unique().tolist())
    kb = []
    
    all_mark = "✅ " if "ALL" in selected_set or not selected_set else ""
    kb.append([InlineKeyboardButton(f"{all_mark}ALL Options", callback_data=f"fl_{dim_name}_ALL")])
    
    row = []
    for opt in opts[:10]:
        mark = "✅ " if opt in selected_set else ""
        row.append(InlineKeyboardButton(f"{mark}{opt}", callback_data=f"fl_{dim_name}_{opt}"))
        if len(row) == 2:
            kb.append(row)
            row = []
    if row:
        kb.append(row)
        
    kb.append([InlineKeyboardButton("➡️ Continue to Timeframe", callback_data="step_timeframe")])
    return InlineKeyboardMarkup(kb)

# ---------------------------------------------------------
# Telegram Handlers (Authentication & Password Enforcement)
# ---------------------------------------------------------
async def start(u: Update, c: ContextTypes.DEFAULT_TYPE):
    user_id = u.effective_user.id
    
    if user_id not in SESSION_CACHE or not SESSION_CACHE[user_id].get("authenticated"):
        c.user_data['login_stage'] = 'AWAITING_EMAIL'
        await u.message.reply_text(
            "🔐 **Data Access Control System**\n\n"
            "Please enter your registered **corporate email address** to begin:"
        )
        return

    c.user_data['filters'] = {}
    user_config = SESSION_CACHE[user_id]
    
    store_info = "All Stores" if user_config['allowed_stores'] == "ALL" else ", ".join(user_config['allowed_stores'])
    await u.message.reply_text(
        f"👋 **Welcome ({user_config['email']})**\n"
        f"🔑 **Role:** `{user_config['role']}`\n"
        f"🏬 **Scope:** `{store_info}`\n"
        f"💰 **Figures:** `Values in ₹ Lacs`\n\n"
        f"Select Primary Dimension:", 
        reply_markup=get_main_menu(), 
        parse_mode="Markdown"
    )

async def forgot_password_command(u: Update, c: ContextTypes.DEFAULT_TYPE):
    user_id = u.effective_user.id
    user_email = SESSION_CACHE.get(user_id, {}).get('email')

    if not user_email:
        await u.message.reply_text(" Please send your registered email ID first using /start.")
        return

    if user_email in USER_PASSWORDS:
        raw_pass = USER_PASSWORDS[user_email]['plain']
        email_sent = send_password_email(user_email, raw_pass)
        
        msg = f"🔑 **Password Recovery**\n\nYour active password is: `{raw_pass}`\n"
        if email_sent:
            msg += f"\n📧 An email containing your password has also been sent to `{user_email}`."
        await u.message.reply_text(msg, parse_mode="Markdown")
    else:
        await u.message.reply_text("No active password found. Please log in using /start.")

async def handle_text_messages(u: Update, c: ContextTypes.DEFAULT_TYPE):
    user_id = u.effective_user.id
    text = u.message.text.strip()
    stage = c.user_data.get('login_stage')

    # If already fully authenticated
    if user_id in SESSION_CACHE and SESSION_CACHE[user_id].get("authenticated"):
        await start(u, c)
        return

    # Stage 1: Check Email
    if stage == 'AWAITING_EMAIL' or "@" in text:
        email = text.lower()
        if email not in AUTHORIZED_USERS:
            await u.message.reply_text("⛔ **Access Denied:** Email address is not authorized. Contact your Operations Lead.")
            return

        c.user_data['pending_email'] = email

        # Generate password if first time
        if email not in USER_PASSWORDS:
            generated_pwd = generate_random_password(8)
            USER_PASSWORDS[email] = {
                "hash": hash_pass(generated_pwd),
                "plain": generated_pwd
            }
            send_password_email(email, generated_pwd)
            
            await u.message.reply_text(
                f"✅ **Email Recognized.**\n\n"
                f"🔑 A new secure password has been generated for your account:\n"
                f"**Password:** `{generated_pwd}`\n\n"
                f"*(Save this password. You can use `/forgotpassword` anytime if you forget it.)*\n\n"
                f"Please reply with this **Password** now to log in:",
                parse_mode="Markdown"
            )
        else:
            await u.message.reply_text(
                f"🔒 **Password Required** for `{email}`:\n\n"
                f"Please enter your password to unlock performance data:",
                parse_mode="Markdown"
            )

        c.user_data['login_stage'] = 'AWAITING_PASSWORD'
        return

    # Stage 2: Validate Password
    if stage == 'AWAITING_PASSWORD':
        email = c.user_data.get('pending_email')
        if not email or email not in USER_PASSWORDS:
            c.user_data['login_stage'] = 'AWAITING_EMAIL'
            await u.message.reply_text("Session expired. Please send your email ID again.")
            return

        entered_hash = hash_pass(text)
        stored_hash = USER_PASSWORDS[email]['hash']

        if entered_hash == stored_hash:
            SESSION_CACHE[user_id] = AUTHORIZED_USERS[email].copy()
            SESSION_CACHE[user_id]['email'] = email
            SESSION_CACHE[user_id]['authenticated'] = True
            c.user_data['login_stage'] = None

            await u.message.reply_text("🔓 **Authentication Successful!** Access Granted.")
            await start(u, c)
        else:
            await u.message.reply_text("❌ **Incorrect Password.** Please try again or use /forgotpassword.")
        return

    # Fallback greeting prompt
    await u.message.reply_text(
        "👋 **Welcome to Analytics Control System**\n\n"
        "Please enter your **registered corporate email address** to continue:"
    )

async def handle_callback(u: Update, c: ContextTypes.DEFAULT_TYPE):
    q = u.callback_query
    await q.answer()
    user_id = q.from_user.id
    
    if user_id not in SESSION_CACHE or not SESSION_CACHE[user_id].get("authenticated"):
        await q.message.reply_text("🔒 **Access Denied.** Please log in using /start.")
        return

    data = q.data
    user_config = SESSION_CACHE[user_id]
    
    if data.startswith("p_"):
        c.user_data['prim'] = data.split("_")[1]
        await q.edit_message_text("🏬 **Select Store Type:**", reply_markup=get_store_type_menu(), parse_mode="Markdown")
        
    elif data.startswith("st_"):
        c.user_data['filters']['Store Type'] = data.split("_")[1]
        dims = [d for d in ['Brand', 'Region', 'Source', 'Session'] if d != c.user_data.get('prim')]
        first_dim = dims[0]
        await q.edit_message_text(f"🔍 **Filter by {first_dim}:**", reply_markup=get_filter_menu(first_dim, set()), parse_mode="Markdown")

    elif data.startswith("fl_"):
        _, dim, val = data.split("_", 2)
        if dim not in c.user_data['filters']:
            c.user_data['filters'][dim] = set()
            
        if val == "ALL":
            c.user_data['filters'][dim] = {"ALL"}
        else:
            c.user_data['filters'][dim].discard("ALL")
            if val in c.user_data['filters'][dim]:
                c.user_data['filters'][dim].remove(val)
            else:
                c.user_data['filters'][dim].add(val)
                
        await q.edit_message_reply_markup(reply_markup=get_filter_menu(dim, c.user_data['filters'][dim]))

    elif data == "step_timeframe":
        await q.edit_message_text("📅 **Select Timeframe:**", reply_markup=get_timeframe_menu(), parse_mode="Markdown")

    elif data.startswith("tf_"):
        tf = data.split("_")[1]
        c.user_data['timeframe'] = tf
        
        df_all, df_eval, months = get_filtered_data(c.user_data['filters'], tf, user_config)
        prim_col = DIM_COL_MAP[c.user_data['prim']]
        piv = generate_pivot(df_all, prim_col, months)
        
        c.user_data['piv'] = piv
        c.user_data['df_eval'] = df_eval
        c.user_data['months'] = months
        
        summary_text = build_telegram_summary(piv, c.user_data['prim'], tf)
        kb = [[InlineKeyboardButton("📄 Export Excel", callback_data="dl_xls"), InlineKeyboardButton("📊 Export PPT", callback_data="dl_ppt")]]
        
        await q.edit_message_text(f"{summary_text}\n\nChoose export format:", reply_markup=InlineKeyboardMarkup(kb), parse_mode="Markdown")
        
    elif data == "dl_xls":
        doc = build_multi_sheet_excel(c.user_data['df_eval'], c.user_data['months'], c.user_data['prim'])
        await c.bot.send_document(q.message.chat_id, doc, filename=f"Analytics_Store_Split_Report.xlsx")
        
    elif data == "dl_ppt":
        doc = build_pptx(c.user_data['df_eval'], c.user_data['months'], c.user_data['prim'], c.user_data['timeframe'])
        await c.bot.send_document(q.message.chat_id, doc, filename=f"Analytics_Presentation.pptx")

async def error_handler(u: object, c: ContextTypes.DEFAULT_TYPE):
    if "Conflict" in str(c.error):
        return
    logger.error("Error encountered:", exc_info=c.error)

if __name__ == '__main__':
    threading.Thread(target=run_flask, daemon=True).start()

    token = os.environ.get("ANALYTICS_BOT_TOKEN") or os.environ.get("TELEGRAM_BOT_TOKEN")
    if not token:
        raise KeyError("Bot token missing in Environment Variables!")

    app = ApplicationBuilder().token(token).build()
    
    app.add_handler(CommandHandler("start", start))
    app.add_handler(CommandHandler("forgotpassword", forgot_password_command))
    app.add_handler(CallbackQueryHandler(handle_callback))
    app.add_handler(MessageHandler(filters.TEXT & ~filters.COMMAND, handle_text_messages))
    app.add_error_handler(error_handler)

    print("📊 Password-Protected Analytics Bot Online...")
    
    app.run_polling(
        drop_pending_updates=True,
        allowed_updates=Update.ALL_TYPES,
        poll_interval=1.0,
        timeout=30
    )
